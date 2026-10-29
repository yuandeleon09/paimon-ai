import io
import os
import re
import zipfile
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
from dataclasses import dataclass, field
from PIL import Image

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

MAX_FILES = 4
MAX_FILE_BYTES = 8 * 1024 * 1024       # 8 MB per file
MAX_TOTAL_MEDIA_BYTES = 12 * 1024 * 1024 # 12 MB combined
MAX_CHARS_PER_FILE = 9000
MAX_TOTAL_CHARS = 28000

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
PDF_EXTENSIONS = {".pdf"}
OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".py", ".js", ".ts",
    ".tsx", ".jsx", ".html", ".css", ".yaml", ".yml",
    ".xml", ".log", ".sql"
}
SUPPORTED_EXTENSIONS = IMAGE_EXTENSIONS | PDF_EXTENSIONS | OFFICE_EXTENSIONS | TEXT_EXTENSIONS

@dataclass
class ProcessedFile:
    filename: str
    extension: str
    size_bytes: int
    category: str
    extracted_text: str = ""
    is_multimodal: bool = False
    mime_type: str = ""
    raw_bytes: bytes = b""
    warnings: List[str] = field(default_factory=list)

@dataclass
class AttachmentEvidencePacket:
    files: List[ProcessedFile] = field(default_factory=list)
    combined_extracted_text: str = ""
    analysis_summary: str = ""
    warnings: List[str] = field(default_factory=list)
    has_attachments: bool = False

    def to_evidence_markdown(self) -> str:
        if not self.has_attachments:
            return ""

        sections = ["## 📎 Attached File Evidence"]
        sections.append(f"**Attached Files ({len(self.files)}):** " + ", ".join([f"`{f.filename}` ({f.size_bytes // 1024} KB)" for f in self.files]))

        if self.warnings:
            sections.append("**Warnings / Extraction Limitations:**")
            for w in self.warnings:
                sections.append(f"- ⚠️ {w}")

        if self.analysis_summary:
            sections.append("### Multimodal Analysis Findings")
            sections.append(self.analysis_summary)

        if self.combined_extracted_text:
            sections.append("### Extracted Text & Data Reference")
            sections.append(self.combined_extracted_text)

        return "\n\n".join(sections)

class AttachmentIngestionService:
    def __init__(self, temp_dir: Optional[Path] = None):
        self.temp_dir = temp_dir or (Path("app_data") / "temp_uploads")
        self.temp_dir.mkdir(parents=True, exist_ok=True)

    def validate_and_process(self, raw_files: List[Tuple[str, bytes]]) -> AttachmentEvidencePacket:
        """
        Validates attachment count, sizes, extension signatures, and extracts content locally.
        Rejects invalid, oversized, or spoofed files early.
        """
        if not raw_files:
            return AttachmentEvidencePacket(has_attachments=False)

        if len(raw_files) > MAX_FILES:
            raise ValueError(f"Too many attachments. Maximum allowed is {MAX_FILES} files (got {len(raw_files)}).")

        total_bytes = sum(len(content) for _, content in raw_files)
        if total_bytes > MAX_TOTAL_MEDIA_BYTES:
            raise ValueError(f"Total attachments size ({total_bytes / (1024*1024):.1f} MB) exceeds maximum limit of 12 MB.")

        processed: List[ProcessedFile] = []
        all_warnings: List[str] = []
        total_extracted_chars = 0

        for filename, data in raw_files:
            size = len(data)
            if size > MAX_FILE_BYTES:
                raise ValueError(f"File '{filename}' ({size / (1024*1024):.1f} MB) exceeds maximum allowed size of 8 MB.")

            ext = Path(filename).suffix.lower()
            if ext not in SUPPORTED_EXTENSIONS:
                raise ValueError(f"Unsupported file type '{ext}' for '{filename}'. Supported: images (PNG/JPG/WEBP), PDF, Word (DOCX), Excel (XLSX), PowerPoint (PPTX), plain text and code.")

            # Validate signatures
            self._verify_signature(filename, ext, data)

            pf = ProcessedFile(
                filename=filename,
                extension=ext,
                size_bytes=size,
                category=self._categorize(ext),
                raw_bytes=data
            )

            # Local extraction by category
            if pf.category == "image":
                pf.is_multimodal = True
                pf.mime_type = "image/png" if ext == ".png" else ("image/jpeg" if ext in (".jpg", ".jpeg") else "image/webp")
            elif pf.category == "pdf":
                pf.is_multimodal = True
                pf.mime_type = "application/pdf"
            elif pf.category == "office":
                text, warns = self._extract_office(filename, ext, data)
                pf.extracted_text = text
                pf.warnings.extend(warns)
                all_warnings.extend(warns)
            elif pf.category == "text":
                text, warns = self._extract_text(filename, data)
                pf.extracted_text = text
                pf.warnings.extend(warns)
                all_warnings.extend(warns)

            # Check character bounds
            if len(pf.extracted_text) > MAX_CHARS_PER_FILE:
                pf.extracted_text = pf.extracted_text[:MAX_CHARS_PER_FILE] + f"\n\n[... Truncated at {MAX_CHARS_PER_FILE} characters ...]"
                w = f"File '{filename}' content was truncated at {MAX_CHARS_PER_FILE} characters to protect model context."
                pf.warnings.append(w)
                all_warnings.append(w)

            total_extracted_chars += len(pf.extracted_text)
            processed.append(pf)

        # Build combined text reference
        combined_text_blocks = []
        for pf in processed:
            if pf.extracted_text:
                combined_text_blocks.append(f"--- File: {pf.filename} ---\n{pf.extracted_text}")

        combined_text = "\n\n".join(combined_text_blocks)
        if len(combined_text) > MAX_TOTAL_CHARS:
            combined_text = combined_text[:MAX_TOTAL_CHARS] + f"\n\n[... Total extracted document text truncated at {MAX_TOTAL_CHARS} characters ...]"
            all_warnings.append(f"Combined document text truncated at {MAX_TOTAL_CHARS} characters.")

        return AttachmentEvidencePacket(
            files=processed,
            combined_extracted_text=combined_text,
            warnings=all_warnings,
            has_attachments=True
        )

    def analyze_multimodal_evidence(self, task: str, packet: AttachmentEvidencePacket, gemini_client: Any) -> str:
        """
        Executes exactly ONE targeted multimodal Gemini call across all attached images and PDFs.
        Extracts visual evidence, errors, and diagrams to share across agents.
        """
        if not packet.has_attachments:
            return ""

        multimodal_files = [f for f in packet.files if f.is_multimodal]
        if not multimodal_files:
            packet.analysis_summary = "No images or PDFs attached. Factual data extracted locally from text and documents."
            return packet.analysis_summary

        # Check if gemini_client is in mock mode or live
        if getattr(gemini_client, "mock_mode", False) or not getattr(gemini_client, "_client", None):
            # Deterministic mock analysis
            summaries = []
            for f in multimodal_files:
                if f.category == "image":
                    summaries.append(f"[{f.filename}]: Visual inspection identified screenshot / diagram related to '{task}'. UI elements and text references parsed.")
                elif f.category == "pdf":
                    summaries.append(f"[{f.filename}]: Document structure inspected. Section headings, tables, and page boundaries identified.")
            packet.analysis_summary = "\n".join(summaries)
            return packet.analysis_summary

        # Live Gemini multimodal call using google.genai types.Part
        try:
            from google.genai import types

            parts = []
            for f in multimodal_files:
                parts.append(types.Part.from_bytes(data=f.raw_bytes, mime_type=f.mime_type))

            prompt = (
                f"You are the Attachment Analyst. Carefully inspect the attached image(s) or PDF(s) in relation to this user task:\n"
                f"Task: {task}\n\n"
                f"Instructions:\n"
                f"1. Extract all key information, visible errors, code snippets, stack traces, diagrams, and tables.\n"
                f"2. Cite filenames and page numbers where identifiable (e.g. '[document.pdf, page 2]').\n"
                f"3. Do not invent page numbers or details if unreadable; explicitly report if any part is blurry or truncated.\n"
                f"4. Provide a concise, structured factual summary for engineering agents."
            )

            contents = [prompt, *parts]
            response = gemini_client._client.models.generate_content(
                model=gemini_client.model,
                contents=contents
            )
            analysis_text = response.text or "Multimodal analysis completed with no visible text detected."
            packet.analysis_summary = analysis_text
            gemini_client.call_count += 1
            return analysis_text
        except Exception as e:
            err = f"Multimodal analysis error: {str(e)}"
            packet.warnings.append(err)
            packet.analysis_summary = f"⚠️ Could not complete visual analysis ({str(e)}). Relying on local document text."
            return packet.analysis_summary

    def _categorize(self, ext: str) -> str:
        if ext in IMAGE_EXTENSIONS:
            return "image"
        if ext in PDF_EXTENSIONS:
            return "pdf"
        if ext in OFFICE_EXTENSIONS:
            return "office"
        return "text"

    def _verify_signature(self, filename: str, ext: str, data: bytes) -> None:
        if ext == ".png":
            if not data.startswith(b"\x89PNG\r\n\x1a\n"):
                raise ValueError(f"Corrupted or spoofed PNG image: '{filename}'. Signature does not match PNG format.")
            try:
                img = Image.open(io.BytesIO(data))
                img.verify()
            except Exception as e:
                raise ValueError(f"Invalid PNG image file '{filename}': {e}")

        elif ext in (".jpg", ".jpeg"):
            if not (data.startswith(b"\xff\xd8\xff") or data[:2] == b"\xff\xd8"):
                raise ValueError(f"Corrupted or spoofed JPEG image: '{filename}'.")
            try:
                img = Image.open(io.BytesIO(data))
                img.verify()
            except Exception as e:
                raise ValueError(f"Invalid JPEG image file '{filename}': {e}")

        elif ext == ".webp":
            if not (data.startswith(b"RIFF") and b"WEBP" in data[:16]):
                raise ValueError(f"Corrupted or spoofed WebP image: '{filename}'.")

        elif ext == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ValueError(f"Invalid or corrupted PDF file: '{filename}'. Header does not match '%PDF-'.")

        elif ext in (".docx", ".xlsx", ".pptx"):
            if not data.startswith(b"PK\x03\x04"):
                raise ValueError(f"Invalid Office archive: '{filename}'. Header is not a valid OpenXML ZIP archive.")
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as zf:
                    bad_file = zf.testzip()
                    if bad_file:
                        raise ValueError(f"Corrupt entry in Office archive: '{bad_file}'")
            except Exception as e:
                raise ValueError(f"Invalid or corrupt Office document: '{filename}'. {str(e)}")

    def _extract_office(self, filename: str, ext: str, data: bytes) -> Tuple[str, List[str]]:
        warns: List[str] = []
        text_lines: List[str] = []

        if ext == ".docx":
            if not Document:
                warns.append("python-docx is not installed. Word extraction unavailable.")
                return "", warns
            try:
                doc = Document(io.BytesIO(data))
                for p in doc.paragraphs:
                    if p.text.strip():
                        text_lines.append(p.text)
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            text_lines.append(" | ".join(row_cells))
            except Exception as e:
                warns.append(f"Failed to extract text from Word document '{filename}': {str(e)}")

        elif ext == ".xlsx":
            if not openpyxl:
                warns.append("openpyxl is not installed. Excel extraction unavailable.")
                return "", warns
            try:
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                for sheet_name in wb.sheetnames[:5]:
                    ws = wb[sheet_name]
                    text_lines.append(f"Sheet: [{sheet_name}]")
                    row_count = 0
                    for row in ws.iter_rows(values_only=True):
                        if row_count >= 100:
                            text_lines.append(f"[... Sheet '{sheet_name}' limited to first 100 rows ...]")
                            break
                        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                        if cells:
                            text_lines.append(" | ".join(cells))
                            row_count += 1
            except Exception as e:
                warns.append(f"Failed to read Excel spreadsheet '{filename}': {str(e)}")

        elif ext == ".pptx":
            if not Presentation:
                warns.append("python-pptx is not installed. PowerPoint extraction unavailable.")
                return "", warns
            try:
                prs = Presentation(io.BytesIO(data))
                for idx, slide in enumerate(prs.slides, 1):
                    slide_texts = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_texts.append(shape.text.strip())
                    if slide_texts:
                        text_lines.append(f"Slide {idx}:\n" + "\n".join(slide_texts))
            except Exception as e:
                warns.append(f"Failed to extract text from PowerPoint '{filename}': {str(e)}")

        return "\n".join(text_lines), warns

    def _extract_text(self, filename: str, data: bytes) -> Tuple[str, List[str]]:
        warns: List[str] = []
        try:
            content = data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                content = data.decode("latin-1")
                warns.append(f"File '{filename}' decoded with latin-1 fallback.")
            except Exception as e:
                warns.append(f"Could not decode text file '{filename}': {str(e)}")
                return "", warns

        return content, warns
